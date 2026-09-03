# Этот файл реализует сервисный уровень: сбор результатов, экспорт в Excel и основной main_with_date.
import os
import re
from typing import Dict, List, Optional, Tuple
import openpyxl
import pandas as pd

from .core import (
    BASE_ERROR_COLUMNS,
    BKI_MAPPING,
    DAILY_GROUP_ALIASES,
    ERROR_CACHE,
    LAST_RUN_STATS,
    MAIN_FILE_CACHE,
    TITLE_CACHE,
    detect_file_type,
    extract_date_from_filename,
    fmt_dt,
    get_xml_cache_stats,
    now_s,
)
from .domain import (
    PairScanResult,
    _resolve_input_dirs,
    detect_group_type,
    find_file_pairs_between_dirs,
    scan_report_response_matches,
    process_pairs_parallel_optimized,
)


def _timer_key_for_group(group_label: str) -> str:
    return "3_2" if group_label == "3-2" else "daily"


def _print_timers(timers: Dict[str, float]) -> None:
    if not timers:
        return
    print("[TIMERS] Detailed timings:")
    for key in sorted(timers):
        try:
            print(f"  {key}: {fmt_dt(float(timers[key]))}")
        except Exception:
            print(f"  {key}: {timers[key]}")


def _add_client_info_ultrafast(errors: List[Dict], main_file: str) -> List[Dict]:
    return errors


def _bki_for_file(path: str) -> str:
    return BKI_MAPPING.get(detect_file_type(path) or "", "")


def _date_matches(path: str, selected_set: set[str]) -> bool:
    if not selected_set:
        return True
    file_date = extract_date_from_filename(os.path.basename(path))
    return not file_date or file_date in selected_set


def filter_scan_result_by_dates(
    scan_result: PairScanResult, selected_dates: List[str], include_unmatched_responses: bool = False
) -> PairScanResult:
    selected_set = {d for d in (selected_dates or []) if d}
    filtered = PairScanResult(duplicate_responses=list(scan_result.duplicate_responses))
    for main_file, response_file in scan_result.pairs:
        if _date_matches(main_file, selected_set):
            filtered.pairs.append((main_file, response_file))
    for main_file in scan_result.reports_without_response:
        if _date_matches(main_file, selected_set):
            filtered.reports_without_response.append(main_file)
    if include_unmatched_responses:
        for response_file in scan_result.responses_without_report:
            response_date = extract_date_from_filename(os.path.basename(response_file))
            if not selected_set or not response_date or response_date in selected_set:
                filtered.responses_without_report.append(response_file)
    return filtered


def _autosize_columns(ws) -> None:
    for column in ws.columns:
        max_len = 0
        letter = column[0].column_letter
        for cell in column:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 80)


def _unique_output_path(output_dir: str, filename: str) -> str:
    path = os.path.join(output_dir or ".", filename)
    if not os.path.exists(path):
        return path
    stem, ext = os.path.splitext(filename)
    counter = 1
    while True:
        candidate = os.path.join(output_dir or ".", f"{stem}_{counter}{ext}")
        if not os.path.exists(candidate):
            return candidate
        counter += 1


def save_missing_receipts_to_excel(missing_reports: List[str], output_dir: str = ".") -> List[str]:
    if not missing_reports:
        return []
    os.makedirs(output_dir or ".", exist_ok=True)
    excel_path = _unique_output_path(output_dir, "missing_receipts.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Нет отбивки"
    ws.append(["Проблема", "Дата", "Отчет", "Путь отчета", "BKI"])
    for report_path in sorted(missing_reports, key=lambda p: (extract_date_from_filename(os.path.basename(p)), p)):
        ws.append(
            [
                "нет отбивки",
                extract_date_from_filename(os.path.basename(report_path)),
                os.path.basename(report_path),
                report_path,
                _bki_for_file(report_path),
            ]
        )
    _autosize_columns(ws)
    wb.save(excel_path)
    return [excel_path]


def save_reconciliation_to_excel(scan_result: PairScanResult, output_dir: str = ".") -> List[str]:
    os.makedirs(output_dir or ".", exist_ok=True)
    excel_path = _unique_output_path(output_dir, "reconciliation.xlsx")
    rows = []
    for report_path in scan_result.reports_without_response:
        rows.append(
            {
                "unknown_date": False,
                "problem": "нет отбивки",
                "date": extract_date_from_filename(os.path.basename(report_path)),
                "file": os.path.basename(report_path),
                "path": report_path,
                "bki": _bki_for_file(report_path),
                "comment": "",
            }
        )
    for response_path in scan_result.responses_without_report:
        response_date = extract_date_from_filename(os.path.basename(response_path))
        rows.append(
            {
                "unknown_date": not bool(response_date),
                "problem": "нет отчета",
                "date": response_date or "дата не определена",
                "file": os.path.basename(response_path),
                "path": response_path,
                "bki": _bki_for_file(response_path),
                "comment": "дата отбивки не определена" if not response_date else "",
            }
        )
    rows.sort(key=lambda r: (not r["unknown_date"], r["date"], r["problem"], r["file"]))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Сверка"
    ws.append(["Проблема", "Дата", "Файл", "Путь", "BKI", "Комментарий"])
    highlight = openpyxl.styles.PatternFill(fill_type="solid", fgColor="FFF2CC")
    for row in rows:
        ws.append([row["problem"], row["date"], row["file"], row["path"], row["bki"], row["comment"]])
        if row["unknown_date"]:
            for cell in ws[ws.max_row]:
                cell.fill = highlight
    _autosize_columns(ws)
    wb.save(excel_path)
    return [excel_path]


def reconcile_reports_receipts(
    local_folder_path: str,
    selected_dates: List[str],
    output_folder: Optional[str] = None,
    reports_folder: Optional[str] = None,
    responses_folder: Optional[str] = None,
) -> List[str]:
    timers: Dict[str, float] = {}
    total_start = now_s()

    def set_timer(name: str, started_at: float) -> None:
        timers[name] = now_s() - started_at

    LAST_RUN_STATS.clear()
    LAST_RUN_STATS.update(
        {
            "mode": "reconciliation",
            "input_dir": local_folder_path,
            "reports_dir": "",
            "responses_dir": "",
            "pairs_found": 0,
            "pairs_after_date_filter": 0,
            "reports_without_response": 0,
            "responses_without_report": 0,
            "reconciliation_file": "",
            "saved_files": 0,
            "timers": timers,
            "note": "",
        }
    )

    output_dir = output_folder or local_folder_path
    os.makedirs(output_dir, exist_ok=True)

    resolve_start = now_s()
    reports_dir, responses_dir = _resolve_input_dirs(local_folder_path, reports_folder, responses_folder)
    set_timer("input.resolve_dirs_s", resolve_start)
    LAST_RUN_STATS["reports_dir"] = reports_dir
    LAST_RUN_STATS["responses_dir"] = responses_dir

    scan_start = now_s()
    scan_result = scan_report_response_matches(reports_dir, responses_dir, include_unmatched_responses=True)
    set_timer("input.scan_s", scan_start)
    LAST_RUN_STATS["pairs_found"] = len(scan_result.pairs)

    filter_start = now_s()
    scan_result = filter_scan_result_by_dates(scan_result, selected_dates, include_unmatched_responses=True)
    set_timer("input.date_filter_s", filter_start)
    LAST_RUN_STATS["pairs_after_date_filter"] = len(scan_result.pairs)
    LAST_RUN_STATS["reports_without_response"] = len(scan_result.reports_without_response)
    LAST_RUN_STATS["responses_without_report"] = len(scan_result.responses_without_report)

    save_start = now_s()
    created_files = save_reconciliation_to_excel(scan_result, output_dir)
    set_timer("save.reconciliation_s", save_start)
    LAST_RUN_STATS["saved_files"] = len(created_files)
    LAST_RUN_STATS["reconciliation_file"] = created_files[0] if created_files else ""
    LAST_RUN_STATS["note"] = "ok"
    set_timer("total_s", total_start)
    _print_timers(timers)
    return created_files


def save_errors_to_excel(
    groups_by_date: Dict[str, List[Tuple[List[Dict], Dict]]], filename_suffix: str, output_dir: str = "."
) -> List[str]:
    created_files: List[str] = []
    all_errors: List[Dict] = []

    for _, date_groups in groups_by_date.items():
        for errors, group_info in date_groups:
            if errors:
                all_errors.extend(_add_client_info_ultrafast(errors, group_info.get("main_file", "")))

    if not all_errors:
        print(f"[SAVE] No errors to export ({filename_suffix}).")
        return []

    suffix_raw = (filename_suffix or "group").strip().lower()
    if suffix_raw in DAILY_GROUP_ALIASES:
        suffix_raw = "daily"
    suffix_safe = re.sub(r"[^A-Za-z0-9._-]+", "_", suffix_raw).strip("._-") or "group"
    excel_filename = f"errors_{suffix_safe}.xlsx"
    excel_path = _unique_output_path(output_dir, excel_filename)
    try:
        t0 = now_s()
        os.makedirs(output_dir or ".", exist_ok=True)
        df = pd.DataFrame(all_errors)

        for col in BASE_ERROR_COLUMNS:
            if col not in df.columns:
                df[col] = ""

        df = df.reindex(columns=BASE_ERROR_COLUMNS)
        df.to_excel(excel_path, index=False, engine="openpyxl")
        elapsed = now_s() - t0
        print(f"[SAVE] '{excel_path}' written in {fmt_dt(elapsed)} ({len(df)} rows)")
        created_files.append(excel_path)
    except Exception as e:
        print(f"[SAVE ERROR] {e}")
        try:
            os.makedirs(output_dir or ".", exist_ok=True)
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = suffix_safe[:31]
            ws.append(BASE_ERROR_COLUMNS)
            wb.save(excel_path)
            created_files.append(excel_path)
            print(f"[SAVE] Empty workbook created: '{excel_path}'")
        except Exception as e2:
            print(f"[SAVE FAIL] {e2}")
    return created_files


def main_with_date(
    local_folder_path: str,
    selected_dates: List[str],
    event_type: str = "all",
    progress_callback=None,
    output_folder: Optional[str] = None,
    reports_folder: Optional[str] = None,
    responses_folder: Optional[str] = None,
    thread_workers: Optional[int] = None,
) -> List[str]:
    total_start = now_s()
    timers: Dict[str, float] = {}

    def set_timer(name: str, started_at: float) -> None:
        timers[name] = now_s() - started_at

    xml_cache_before = get_xml_cache_stats()
    LAST_RUN_STATS.clear()
    LAST_RUN_STATS.update(
        {
            "input_dir": local_folder_path,
            "reports_dir": "",
            "responses_dir": "",
            "pairs_found": 0,
            "pairs_after_date_filter": 0,
            "reports_without_response": 0,
            "responses_without_report": 0,
            "missing_receipts_file": "",
            "reconciliation_file": "",
            "mode": "errors",
            "group_3_2_pairs": 0,
            "group_daily_pairs": 0,
            "event_type": event_type,
            "errors_found": 0,
            "saved_files": 0,
            "xml_cache_hits": 0,
            "xml_cache_misses": 0,
            "xml_cache_invalidated": 0,
            "xml_cache_size": 0,
            "thread_workers": thread_workers if thread_workers is not None else "auto",
            "timers": timers,
            "note": "",
        }
    )
    setup_start = now_s()
    MAIN_FILE_CACHE.clear()
    TITLE_CACHE.clear()
    ERROR_CACHE.clear()
    set_timer("setup.clear_caches_s", setup_start)

    if progress_callback:
        progress_callback(5, "Scanning input files")

    output_start = now_s()
    output_dir = output_folder or local_folder_path
    os.makedirs(output_dir, exist_ok=True)
    set_timer("setup.output_dir_s", output_start)

    resolve_start = now_s()
    reports_dir, responses_dir = _resolve_input_dirs(local_folder_path, reports_folder, responses_folder)
    set_timer("input.resolve_dirs_s", resolve_start)
    LAST_RUN_STATS["reports_dir"] = reports_dir
    LAST_RUN_STATS["responses_dir"] = responses_dir
    print(f"[INPUT] Resolved dirs: reports='{reports_dir}', responses='{responses_dir}'")

    pair_start = now_s()
    scan_result = scan_report_response_matches(reports_dir, responses_dir, include_unmatched_responses=False)
    set_timer("input.find_pairs_s", pair_start)
    LAST_RUN_STATS["pairs_found"] = len(scan_result.pairs)

    filter_start = now_s()
    scan_result = filter_scan_result_by_dates(scan_result, selected_dates, include_unmatched_responses=False)
    file_pairs = scan_result.pairs
    set_timer("input.date_filter_s", filter_start)
    LAST_RUN_STATS["pairs_after_date_filter"] = len(file_pairs)
    LAST_RUN_STATS["reports_without_response"] = len(scan_result.reports_without_response)

    created_files: List[str] = []
    missing_save_start = now_s()
    missing_files = save_missing_receipts_to_excel(scan_result.reports_without_response, output_dir=output_dir)
    set_timer("input.save_missing_receipts_s", missing_save_start)
    created_files.extend(missing_files)
    LAST_RUN_STATS["missing_receipts_file"] = missing_files[0] if missing_files else ""

    if not file_pairs:
        if created_files:
            LAST_RUN_STATS["saved_files"] = len(created_files)
            LAST_RUN_STATS["note"] = "missing_receipts_only"
            set_timer("total_s", total_start)
            _print_timers(timers)
            if progress_callback:
                progress_callback(100, f"Completed. Files created: {len(created_files)}")
            return created_files
        LAST_RUN_STATS["note"] = "pairs_not_found"
        raise ValueError("No report/response pairs found. Check file names and reports/responses folders.")

    group_start = now_s()
    daily_group_key = detect_group_type("__daily_group_key__")
    pairs_3_2: List[Tuple[str, str]] = []
    pairs_daily: List[Tuple[str, str]] = []
    for main_file, response_file in file_pairs:
        if detect_group_type(main_file) == "3-2":
            pairs_3_2.append((main_file, response_file))
        else:
            pairs_daily.append((main_file, response_file))
    set_timer("input.group_detection_s", group_start)
    LAST_RUN_STATS["group_3_2_pairs"] = len(pairs_3_2)
    LAST_RUN_STATS["group_daily_pairs"] = len(pairs_daily)
    print(f"[GROUPS] 3-2={len(pairs_3_2)}, daily={len(pairs_daily)}")

    normalized_event_type = (event_type or "all").lower().strip()
    groups_to_process: List[Tuple[str, List[Tuple[str, str]], int, int]] = []
    if normalized_event_type == "3-2":
        groups_to_process.append(("3-2", pairs_3_2, 15, 90))
    elif normalized_event_type in DAILY_GROUP_ALIASES:
        groups_to_process.append((daily_group_key, pairs_daily, 15, 90))
    else:
        if pairs_3_2:
            groups_to_process.append(("3-2", pairs_3_2, 10, 55))
        if pairs_daily:
            groups_to_process.append((daily_group_key, pairs_daily, 55, 90))

    if not groups_to_process:
        raise ValueError("No pairs available for selected event type.")

    total_errors_found = 0
    for group_label, pairs, start_progress, end_progress in groups_to_process:
        if not pairs:
            continue
        group_key = _timer_key_for_group(group_label)
        group_start = now_s()
        process_start = now_s()
        grouped_errors = process_pairs_parallel_optimized(
            pairs,
            group_label,
            progress_callback=progress_callback,
            start_progress=start_progress,
            end_progress=end_progress,
            thread_workers=thread_workers,
            timers=timers,
            timer_prefix=f"group.{group_key}.threadpool",
        )
        set_timer(f"group.{group_key}.process_pairs_s", process_start)
        count_start = now_s()
        group_errors_count = sum(len(errors or []) for date_groups in grouped_errors.values() for errors, _ in date_groups)
        set_timer(f"group.{group_key}.count_errors_s", count_start)
        total_errors_found += group_errors_count
        print(f"[RESULT] group='{group_label}' pairs={len(pairs)} errors={group_errors_count}")
        save_start = now_s()
        created_files.extend(save_errors_to_excel(grouped_errors, group_label, output_dir=output_dir))
        set_timer(f"group.{group_key}.save_excel_s", save_start)
        set_timer(f"group.{group_key}.total_s", group_start)

    finalize_start = now_s()
    LAST_RUN_STATS["errors_found"] = total_errors_found
    LAST_RUN_STATS["saved_files"] = len(created_files)
    xml_cache_after = get_xml_cache_stats()
    LAST_RUN_STATS["xml_cache_hits"] = max(0, xml_cache_after["hits"] - xml_cache_before["hits"])
    LAST_RUN_STATS["xml_cache_misses"] = max(0, xml_cache_after["misses"] - xml_cache_before["misses"])
    LAST_RUN_STATS["xml_cache_invalidated"] = max(0, xml_cache_after["invalidated"] - xml_cache_before["invalidated"])
    LAST_RUN_STATS["xml_cache_size"] = xml_cache_after["size"]
    set_timer("finalize.stats_s", finalize_start)
    if not created_files and total_errors_found == 0:
        LAST_RUN_STATS["note"] = "no_errors_in_receipts"
    elif not created_files:
        LAST_RUN_STATS["note"] = "save_failed_or_empty"
    else:
        LAST_RUN_STATS["note"] = "ok"
    print(
        f"[SUMMARY] pairs={LAST_RUN_STATS['pairs_found']} "
        f"filtered={LAST_RUN_STATS['pairs_after_date_filter']} "
        f"errors={total_errors_found} files={len(created_files)} "
        f"cache(h={LAST_RUN_STATS['xml_cache_hits']},m={LAST_RUN_STATS['xml_cache_misses']},"
        f"i={LAST_RUN_STATS['xml_cache_invalidated']},size={LAST_RUN_STATS['xml_cache_size']}) "
        f"note={LAST_RUN_STATS['note']}"
    )
    set_timer("total_s", total_start)
    _print_timers(timers)

    if progress_callback:
        progress_callback(100, f"Completed. Files created: {len(created_files)}")
    return created_files
