import os
import tempfile
import unittest

import openpyxl

from check_receipts.domain import scan_report_response_matches
from check_receipts.service import (
    filter_scan_result_by_dates,
    main_with_date,
    reconcile_reports_receipts,
    save_reconciliation_to_excel,
)
from orchestration.runner import run_processing
from orchestration.sources import expand_creditline_sources
from orchestration.state import processing_status


def touch(path):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        f.write("<xml/>")


class ReconciliationTests(unittest.TestCase):
    def test_standard_scan_does_not_report_unmatched_responses(self):
        with tempfile.TemporaryDirectory() as tmp:
            reports = os.path.join(tmp, "Reports")
            receipts = os.path.join(tmp, "Receipts")
            touch(os.path.join(reports, "BD0_20260510.xml"))
            touch(os.path.join(reports, "BD0_20260511.xml"))
            touch(os.path.join(receipts, "BD0_20260510_TICKET2.xml"))
            touch(os.path.join(receipts, "BD0_20260512_TICKET2.xml"))

            result = scan_report_response_matches(reports, receipts, include_unmatched_responses=False)

            self.assertEqual(1, len(result.pairs))
            self.assertEqual(["BD0_20260511.xml"], [os.path.basename(p) for p in result.reports_without_response])
            self.assertEqual([], result.responses_without_report)

    def test_reconciliation_scan_reports_both_missing_sides(self):
        with tempfile.TemporaryDirectory() as tmp:
            reports = os.path.join(tmp, "Reports")
            receipts = os.path.join(tmp, "Receipts")
            touch(os.path.join(reports, "BD0_20260510.xml"))
            touch(os.path.join(reports, "BD0_20260511.xml"))
            touch(os.path.join(receipts, "BD0_20260510_TICKET2.xml"))
            touch(os.path.join(receipts, "BD0_20260512_TICKET2.xml"))

            result = scan_report_response_matches(reports, receipts, include_unmatched_responses=True)

            self.assertEqual(1, len(result.pairs))
            self.assertEqual(["BD0_20260511.xml"], [os.path.basename(p) for p in result.reports_without_response])
            self.assertEqual(["BD0_20260512_TICKET2.xml"], [os.path.basename(p) for p in result.responses_without_report])

    def test_reconciliation_excel_keeps_unknown_date_receipts_first_and_highlighted(self):
        with tempfile.TemporaryDirectory() as tmp:
            reports = os.path.join(tmp, "Reports")
            receipts = os.path.join(tmp, "Receipts")
            output = os.path.join(tmp, "out")
            touch(os.path.join(reports, "BD0_20260511.xml"))
            touch(os.path.join(receipts, "BD0_WITHOUT_DATE_TICKET2.xml"))
            touch(os.path.join(receipts, "BD0_20260512_TICKET2.xml"))

            scan = scan_report_response_matches(reports, receipts, include_unmatched_responses=True)
            filtered = filter_scan_result_by_dates(scan, ["2026-05-11"], include_unmatched_responses=True)
            files = save_reconciliation_to_excel(filtered, output)

            self.assertEqual([os.path.join(output, "reconciliation.xlsx")], files)
            wb = openpyxl.load_workbook(files[0])
            ws = wb["Сверка"]
            self.assertEqual("BD0_WITHOUT_DATE_TICKET2.xml", ws["C2"].value)
            self.assertEqual("дата не определена", ws["B2"].value)
            self.assertIsNotNone(ws["A2"].fill.fgColor.rgb)
            self.assertEqual(3, ws.max_row)

    def test_standard_run_succeeds_with_missing_receipts_when_no_pairs_found(self):
        with tempfile.TemporaryDirectory() as tmp:
            reports = os.path.join(tmp, "Reports")
            receipts = os.path.join(tmp, "Receipts")
            output = os.path.join(tmp, "out")
            touch(os.path.join(reports, "BD0_20260511.xml"))
            os.makedirs(receipts, exist_ok=True)

            files = main_with_date(
                tmp,
                ["2026-05-11"],
                output_folder=output,
                reports_folder=reports,
                responses_folder=receipts,
            )

            self.assertEqual([os.path.join(output, "missing_receipts.xlsx")], files)

    def test_reconciliation_mode_creates_reconciliation_workbook(self):
        with tempfile.TemporaryDirectory() as tmp:
            reports = os.path.join(tmp, "Reports")
            receipts = os.path.join(tmp, "Receipts")
            output = os.path.join(tmp, "out")
            touch(os.path.join(reports, "BD0_20260511.xml"))
            touch(os.path.join(receipts, "BD0_20260512_TICKET2.xml"))

            files = reconcile_reports_receipts(
                tmp,
                ["2026-05-11"],
                output_folder=output,
                reports_folder=reports,
                responses_folder=receipts,
            )

            self.assertEqual([os.path.join(output, "reconciliation.xlsx")], files)

    def test_reconciliation_does_not_overwrite_existing_workbook(self):
        with tempfile.TemporaryDirectory() as tmp:
            reports = os.path.join(tmp, "Reports")
            receipts = os.path.join(tmp, "Receipts")
            output = os.path.join(tmp, "out")
            touch(os.path.join(reports, "BD0_20260511.xml"))
            os.makedirs(receipts, exist_ok=True)
            os.makedirs(output, exist_ok=True)
            existing = os.path.join(output, "reconciliation.xlsx")
            touch(existing)

            files = reconcile_reports_receipts(
                tmp,
                ["2026-05-11"],
                output_folder=output,
                reports_folder=reports,
                responses_folder=receipts,
            )

            self.assertEqual([os.path.join(output, "reconciliation_1.xlsx")], files)
            self.assertTrue(os.path.exists(existing))

    def test_creditline_reconciliation_combines_bki_sources_in_one_workbook(self):
        with tempfile.TemporaryDirectory() as tmp:
            creditline = os.path.join(tmp, "CreditLine")
            output = os.path.join(tmp, "out")

            eq_reports = os.path.join(creditline, "EquifaxNew", "Reports")
            eq_receipts = os.path.join(creditline, "EquifaxNew", "Receipts")
            ucb_reports = os.path.join(creditline, "UCBNew", "Reports")
            ucb_receipts = os.path.join(creditline, "UCBNew", "Receipts")
            touch(os.path.join(eq_reports, "0XY_FCH_20260511.xml"))
            os.makedirs(eq_receipts, exist_ok=True)
            touch(os.path.join(ucb_reports, "CHP_20260511.xml"))
            os.makedirs(ucb_receipts, exist_ok=True)

            sources = expand_creditline_sources([creditline])
            run_processing(
                sources,
                output,
                ["2026-05-11"],
                "all",
                mode="reconciliation",
            )

            self.assertFalse(processing_status["has_error"])
            self.assertEqual(["reconciliation.xlsx"], processing_status["result_files"])
            workbook_path = os.path.join(output, "reconciliation.xlsx")
            wb = openpyxl.load_workbook(workbook_path)
            ws = wb.active
            files = {ws["C2"].value, ws["C3"].value}
            self.assertEqual({"0XY_FCH_20260511.xml", "CHP_20260511.xml"}, files)


if __name__ == "__main__":
    unittest.main()
